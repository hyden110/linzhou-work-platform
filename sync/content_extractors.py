"""Read-only, bounded content extraction for local knowledge files."""

from __future__ import annotations

import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable, Optional, Tuple, Union

DEFAULT_MAX_SEGMENT_CHARACTERS = 2_000
DEFAULT_MAX_TOTAL_CHARACTERS = 20_000
SUPPORTED_EXTENSIONS = frozenset({".md", ".pdf", ".pptx", ".docx", ".xlsx"})
UNSAFE_LEGACY_EXTENSIONS = frozenset({".doc", ".xls", ".ppt"})
_MARKDOWN_HEADING = re.compile(r"^(#{1,6})[ \t]+(.+?)\s*#*\s*$")


@dataclass(frozen=True)
class ContentSegment:
    """A bounded piece of text and its location in the source document."""

    title: str
    locator_kind: str
    locator: str
    text: str


@dataclass(frozen=True)
class ExtractionResult:
    """Immutable outcome of one local file extraction attempt."""

    status: str
    file_type: str
    title: str
    segments: Tuple[ContentSegment, ...] = ()
    total_characters: int = 0
    truncated: bool = False
    error_code: Optional[str] = None
    error_message: Optional[str] = None


@dataclass(frozen=True)
class _RawSegment:
    title: str
    locator_kind: str
    locator: str
    text: str


@dataclass(frozen=True)
class _ExtractedContent:
    title: str
    segments: Tuple[_RawSegment, ...]


class _ExtractionFailure(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.safe_message = message


def _normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


def _display_title(value: object, fallback: str) -> str:
    normalized = _normalize_text(value)
    return normalized.split("\n", 1)[0] if normalized else fallback


def _bounded_segments(
    raw_segments: Iterable[_RawSegment],
    max_segment_characters: int,
    max_total_characters: int,
) -> tuple[Tuple[ContentSegment, ...], bool]:
    source_segments = tuple(raw_segments)
    output = []
    consumed = 0
    truncated = False
    for raw in source_segments:
        text = _normalize_text(raw.text)
        offset = 0
        while offset < len(text) and consumed < max_total_characters:
            chunk_size = min(max_segment_characters, max_total_characters - consumed)
            chunk = text[offset : offset + chunk_size]
            output.append(ContentSegment(raw.title, raw.locator_kind, raw.locator, chunk))
            consumed += len(chunk)
            offset += len(chunk)
        if offset < len(text):
            truncated = True
            break
    if not truncated and consumed >= max_total_characters:
        truncated = any(_normalize_text(item.text) for item in source_segments[len(output) :])
    return tuple(output), truncated


def _flush_markdown_section(
    segments: list[_RawSegment],
    heading_path: list[str],
    body: list[str],
    fallback_title: str,
) -> None:
    text = _normalize_text("\n".join(body))
    if not text:
        return
    title = heading_path[-1] if heading_path else fallback_title
    locator = " / ".join(heading_path) if heading_path else "Document"
    segments.append(_RawSegment(title, "section", locator, text))


def _extract_markdown(path: Path) -> _ExtractedContent:
    lines = path.read_text(encoding="utf-8").splitlines()
    segments: list[_RawSegment] = []
    heading_path: list[str] = []
    body: list[str] = []
    document_title = path.stem
    found_title = False
    for line in lines:
        match = _MARKDOWN_HEADING.match(line)
        if not match:
            body.append(line)
            continue
        _flush_markdown_section(segments, heading_path, body, path.stem)
        body = []
        level = len(match.group(1))
        heading = _display_title(match.group(2), path.stem)
        if level == 1 and not found_title:
            document_title = heading
            found_title = True
        heading_path = heading_path[: level - 1] + [heading]
    _flush_markdown_section(segments, heading_path, body, path.stem)
    return _ExtractedContent(document_title, tuple(segments))


def _iter_docx_blocks(document: object) -> Iterable[object]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _docx_table_text(table: object) -> str:
    rows = []
    for row in table.rows:
        values = [_normalize_text(cell.text).replace("\n", " ") for cell in row.cells]
        while values and not values[-1]:
            values.pop()
        if values:
            rows.append("\t".join(values))
    return "\n".join(rows)


def _extract_docx(path: Path) -> _ExtractedContent:
    from docx import Document
    from docx.table import Table

    document = Document(path)
    title = _display_title(document.core_properties.title, path.stem)
    segments: list[_RawSegment] = []
    heading_path: list[str] = []
    body: list[str] = []
    for block in _iter_docx_blocks(document):
        if isinstance(block, Table):
            table_text = _docx_table_text(block)
            if table_text:
                body.append(table_text)
            continue
        text = _normalize_text(block.text)
        style_name = str(block.style.name or "")
        match = re.match(r"Heading\s+(\d+)$", style_name, re.IGNORECASE)
        if match and text:
            _flush_markdown_section(segments, heading_path, body, title)
            body = []
            level = max(1, int(match.group(1)))
            heading_path = heading_path[: level - 1] + [text]
        elif text:
            body.append(text)
    _flush_markdown_section(segments, heading_path, body, title)
    return _ExtractedContent(title, tuple(segments))


def _shape_text(shape: object) -> str:
    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:
            values = [_normalize_text(cell.text).replace("\n", " ") for cell in row.cells]
            while values and not values[-1]:
                values.pop()
            if values:
                rows.append("\t".join(values))
        return "\n".join(rows)
    return _normalize_text(shape.text) if getattr(shape, "has_text_frame", False) else ""


def _extract_pptx(path: Path) -> _ExtractedContent:
    from pptx import Presentation

    presentation = Presentation(path)
    title = _display_title(presentation.core_properties.title, path.stem)
    segments = []
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        slide_title = _display_title(title_shape.text if title_shape else None, f"Slide {number}")
        texts = []
        for shape in slide.shapes:
            if title_shape is not None and shape == title_shape:
                continue
            text = _shape_text(shape)
            if text:
                texts.append(text)
        text = _normalize_text("\n".join(texts))
        if text:
            segments.append(_RawSegment(slide_title, "slide", str(number), text))
    return _ExtractedContent(title, tuple(segments))


def _extract_pdf(path: Path) -> _ExtractedContent:
    from pypdf import PdfReader

    reader = PdfReader(path, strict=False)
    if reader.is_encrypted:
        raise _ExtractionFailure("encrypted_document", "Encrypted PDF files are not supported.")
    metadata = reader.metadata
    title = _display_title(metadata.title if metadata else None, path.stem)
    segments = []
    for number, page in enumerate(reader.pages, start=1):
        text = _normalize_text(page.extract_text())
        if text:
            segments.append(_RawSegment(f"Page {number}", "page", str(number), text))
    return _ExtractedContent(title, tuple(segments))


def _cell_text(value: object) -> str:
    return _normalize_text(value).replace("\n", " ")


def _extract_xlsx(path: Path) -> _ExtractedContent:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=path, read_only=True, data_only=True, keep_links=False)
    try:
        title = _display_title(workbook.properties.title, path.stem)
        segments = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = [_cell_text(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if values:
                    rows.append("\t".join(values))
            text = _normalize_text("\n".join(rows))
            if text:
                segments.append(_RawSegment(sheet.title, "sheet", sheet.title, text))
        return _ExtractedContent(title, tuple(segments))
    finally:
        workbook.close()


_EXTRACTORS: dict[str, Callable[[Path], _ExtractedContent]] = {
    ".md": _extract_markdown,
    ".pdf": _extract_pdf,
    ".pptx": _extract_pptx,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
}


def _validate_limit(name: str, value: int) -> None:
    if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
        raise ValueError(f"{name} must be a positive integer")


def _failure_result(path: Path, code: str, message: str, status: str = "error") -> ExtractionResult:
    return ExtractionResult(
        status=status,
        file_type=path.suffix.lower(),
        title=path.stem,
        error_code=code,
        error_message=message,
    )


def extract_content(
    path: Union[str, os.PathLike[str]],
    *,
    max_segment_characters: int = DEFAULT_MAX_SEGMENT_CHARACTERS,
    max_total_characters: int = DEFAULT_MAX_TOTAL_CHARACTERS,
) -> ExtractionResult:
    """Extract bounded text without modifying or copying the source file."""

    if not isinstance(path, (str, os.PathLike)):
        raise TypeError("path must be a string or path-like object")
    _validate_limit("max_segment_characters", max_segment_characters)
    _validate_limit("max_total_characters", max_total_characters)
    source = Path(path)
    suffix = source.suffix.lower()
    if not source.exists():
        return _failure_result(source, "file_not_found", "Source file does not exist.")
    if not source.is_file():
        return _failure_result(source, "not_a_file", "Source path is not a regular file.")
    if suffix in UNSAFE_LEGACY_EXTENSIONS:
        return _failure_result(
            source,
            "unsafe_legacy_format",
            "Legacy Office formats are unsupported; no conversion was attempted.",
            status="unsupported",
        )
    if suffix not in SUPPORTED_EXTENSIONS:
        return _failure_result(source, "unsupported_format", "File format is not supported.", status="unsupported")
    try:
        extracted = _EXTRACTORS[suffix](source)
        segments, truncated = _bounded_segments(
            extracted.segments,
            max_segment_characters,
            max_total_characters,
        )
        return ExtractionResult(
            status="extracted",
            file_type=suffix,
            title=extracted.title,
            segments=segments,
            total_characters=sum(len(segment.text) for segment in segments),
            truncated=truncated,
        )
    except _ExtractionFailure as error:
        return _failure_result(source, error.code, error.safe_message)
    except ModuleNotFoundError as error:
        return _failure_result(
            source,
            "dependency_unavailable",
            f"Required parser dependency is unavailable ({error.name or 'unknown'}).",
        )
    except Exception as error:
        return _failure_result(
            source,
            "parse_error",
            f"Unable to parse {suffix} content ({type(error).__name__}).",
        )
