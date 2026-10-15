import tempfile
import unittest
from dataclasses import FrozenInstanceError
from pathlib import Path

from sync import content_extractors


class ContentExtractorsTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_extracts_markdown_sections_into_immutable_bounded_segments(self) -> None:
        path = self.root / "guide.md"
        path.write_text(
            "# Product Guide\n\nOverview text.\n\n## Setup\n\nabcdefghijk\n",
            encoding="utf-8",
        )

        result = content_extractors.extract_content(
            path,
            max_segment_characters=5,
            max_total_characters=14,
        )

        self.assertEqual(result.status, "extracted")
        self.assertEqual(result.file_type, ".md")
        self.assertEqual(result.title, "Product Guide")
        self.assertEqual(result.total_characters, 14)
        self.assertTrue(result.truncated)
        self.assertTrue(result.segments)
        self.assertTrue(all(len(segment.text) <= 5 for segment in result.segments))
        self.assertEqual(result.segments[0].locator_kind, "section")
        self.assertEqual(result.segments[0].locator, "Product Guide")
        self.assertIsInstance(result.segments, tuple)
        with self.assertRaises(FrozenInstanceError):
            result.status = "changed"
        with self.assertRaises(FrozenInstanceError):
            result.segments[0].text = "changed"

    def test_marks_result_truncated_when_limit_ends_between_sections(self) -> None:
        path = self.root / "bounded.md"
        path.write_text("# First\n1234567890\n## Second\nmore", encoding="utf-8")

        result = content_extractors.extract_content(
            path,
            max_segment_characters=4,
            max_total_characters=10,
        )

        self.assertEqual(result.total_characters, 10)
        self.assertTrue(result.truncated)
        self.assertTrue(all(len(segment.text) <= 4 for segment in result.segments))

    def test_extracts_pdf_pages(self) -> None:
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        path = self.root / "brief.pdf"
        writer = PdfWriter()
        writer.add_metadata({"/Title": "PDF Brief"})
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_reference = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_reference}
                )
            }
        )
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 72 720 Td (PDF page text) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        with path.open("wb") as stream:
            writer.write(stream)

        result = content_extractors.extract_content(path)

        self.assertEqual(result.status, "extracted")
        self.assertEqual(result.title, "PDF Brief")
        self.assertEqual(len(result.segments), 1)
        self.assertEqual(result.segments[0].locator_kind, "page")
        self.assertEqual(result.segments[0].locator, "1")
        self.assertIn("PDF page text", result.segments[0].text)

    def test_extracts_presentation_slides(self) -> None:
        from pptx import Presentation

        path = self.root / "deck.pptx"
        presentation = Presentation()
        presentation.core_properties.title = "Quarterly Deck"
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Results"
        slide.placeholders[1].text = "Revenue increased"
        presentation.save(path)

        result = content_extractors.extract_content(path)

        self.assertEqual(result.status, "extracted")
        self.assertEqual(result.title, "Quarterly Deck")
        self.assertEqual(result.segments[0].title, "Results")
        self.assertEqual(result.segments[0].locator_kind, "slide")
        self.assertEqual(result.segments[0].locator, "1")
        self.assertIn("Revenue increased", result.segments[0].text)

    def test_extracts_word_sections(self) -> None:
        from docx import Document

        path = self.root / "proposal.docx"
        document = Document()
        document.core_properties.title = "Customer Proposal"
        document.add_heading("Context", level=1)
        document.add_paragraph("Current situation")
        document.add_heading("Plan", level=2)
        document.add_paragraph("Recommended action")
        document.save(path)

        result = content_extractors.extract_content(path)

        self.assertEqual(result.status, "extracted")
        self.assertEqual(result.title, "Customer Proposal")
        self.assertEqual(
            [(segment.locator_kind, segment.locator, segment.title) for segment in result.segments],
            [("section", "Context", "Context"), ("section", "Context / Plan", "Plan")],
        )
        self.assertIn("Current situation", result.segments[0].text)
        self.assertIn("Recommended action", result.segments[1].text)

    def test_extracts_spreadsheet_sheets_in_read_only_data_only_mode(self) -> None:
        from openpyxl import Workbook

        path = self.root / "metrics.xlsx"
        workbook = Workbook()
        workbook.properties.title = "Metrics Workbook"
        sheet = workbook.active
        sheet.title = "Summary"
        sheet.append(["Metric", "Value"])
        sheet.append(["Revenue", 42])
        workbook.create_sheet("Empty")
        workbook.save(path)

        result = content_extractors.extract_content(path)

        self.assertEqual(result.status, "extracted")
        self.assertEqual(result.title, "Metrics Workbook")
        self.assertEqual(len(result.segments), 1)
        self.assertEqual(result.segments[0].locator_kind, "sheet")
        self.assertEqual(result.segments[0].locator, "Summary")
        self.assertEqual(result.segments[0].title, "Summary")
        self.assertIn("Metric\tValue", result.segments[0].text)
        self.assertIn("Revenue\t42", result.segments[0].text)

    def test_returns_unsupported_without_reading_legacy_office_files(self) -> None:
        for suffix in (".doc", ".xls"):
            with self.subTest(suffix=suffix):
                path = self.root / f"legacy{suffix}"
                path.write_bytes(b"legacy source bytes")

                result = content_extractors.extract_content(path)

                self.assertEqual(result.status, "unsupported")
                self.assertEqual(result.error_code, "unsafe_legacy_format")
                self.assertEqual(result.segments, ())
                self.assertEqual(result.total_characters, 0)

    def test_validates_boundaries_and_converts_single_file_errors(self) -> None:
        missing = content_extractors.extract_content(self.root / "missing.pdf")
        self.assertEqual(missing.status, "error")
        self.assertEqual(missing.error_code, "file_not_found")

        unsupported_path = self.root / "notes.txt"
        unsupported_path.write_text("not supported", encoding="utf-8")
        unsupported = content_extractors.extract_content(unsupported_path)
        self.assertEqual(unsupported.status, "unsupported")
        self.assertEqual(unsupported.error_code, "unsupported_format")

        broken_path = self.root / "broken.docx"
        broken_path.write_bytes(b"not a valid docx")
        broken = content_extractors.extract_content(broken_path)
        self.assertEqual(broken.status, "error")
        self.assertEqual(broken.error_code, "parse_error")
        self.assertEqual(broken.segments, ())

        directory = content_extractors.extract_content(self.root)
        self.assertEqual(directory.status, "error")
        self.assertEqual(directory.error_code, "not_a_file")

        with self.assertRaises(TypeError):
            content_extractors.extract_content(123)
        with self.assertRaises(ValueError):
            content_extractors.extract_content(unsupported_path, max_segment_characters=0)
        with self.assertRaises(ValueError):
            content_extractors.extract_content(unsupported_path, max_total_characters=-1)


if __name__ == "__main__":
    unittest.main()
